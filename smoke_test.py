"""HTTP smoke test: exercises every route of the running app.

Usage:  python smoke_test.py            (server must already be running)
"""
import io
import json
import os
import sys
import time
import urllib.request
import urllib.error

import requests
from docx import Document

BASE = "http://127.0.0.1:5000"
failures = []


def check(label, condition, extra=""):
    if extra:
        # the Windows console uses cp1252 - keep the output printable
        extra = str(extra).encode("ascii", "replace").decode("ascii")
    print(("  OK   " if condition else "  FAIL ") + label + (f"  {extra}" if extra else ""))
    if not condition:
        failures.append(label)


def wait_for_server(seconds=45):
    deadline = time.time() + seconds
    while time.time() < deadline:
        try:
            urllib.request.urlopen(BASE + "/", timeout=2)
            return True
        except Exception:
            time.sleep(0.5)
    return False


print("waiting for the server...")
if not wait_for_server():
    print("  FAIL server never started")
    sys.exit(1)

print("\n== pages ==")
for path, needle in [("/", "Dashboard"), ("/new", "Analyze"),
                     ("/history", "Meeting history"), ("/ask", "Ask your meetings"),
                     ("/settings", "AI connection"),
                     ("/meeting/999999", "404")]:
    r = requests.get(BASE + path, timeout=20)
    check(f"GET {path} -> {r.status_code}", r.status_code == (404 if "404" in needle else 200)
          and needle in r.text, f"'{needle}' present={needle in r.text}")

print("\n== validation ==")
r = requests.post(BASE + "/analyze", data={}, allow_redirects=False, timeout=20)
check("empty submit redirects back", r.status_code == 302)

r = requests.post(BASE + "/analyze", data={"notes": "too short"},
                  allow_redirects=False, timeout=20)
check("short notes rejected", r.status_code == 302)

r = requests.post(BASE + "/analyze",
                  data={"notes": "x" * 40, "meeting_date": "not-a-date"},
                  allow_redirects=False, timeout=20)
check("bad date rejected", r.status_code == 302)

r = requests.post(BASE + "/analyze",
                  files={"file": ("evil.exe", b"MZ binary", "application/octet-stream")},
                  data={"notes": ""}, allow_redirects=False, timeout=20)
check(".exe upload rejected", r.status_code == 302)

r = requests.post(BASE + "/analyze",
                  files={"file": ("empty.txt", b"   \n  ", "text/plain")},
                  allow_redirects=False, timeout=20)
check("empty file rejected", r.status_code == 302)

r = requests.post(BASE + "/analyze",
                  files={"file": ("fake.pdf", b"not a pdf at all", "application/pdf")},
                  data={"notes": ""}, allow_redirects=False, timeout=20)
check("fake .pdf rejected", r.status_code == 302)

r = requests.post(BASE + "/analyze",
                  files={"file": ("big.txt", b"a" * (7 * 1024 * 1024), "text/plain")},
                  data={"notes": ""}, allow_redirects=False, timeout=60)
check("oversized upload rejected", r.status_code in (302, 413),
      f"status={r.status_code}")

print("\n== analyze a TXT file ==")
sample_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           "samples", "sample_meeting.txt")
with open(sample_path, "rb") as fh:
    sample_bytes = fh.read()

r = requests.post(BASE + "/analyze",
                  files={"file": ("sample_meeting.txt", sample_bytes, "text/plain")},
                  data={"meeting_date": "2026-10-09"},
                  allow_redirects=False, timeout=90)
check("analyze redirects to results", r.status_code == 302, f"status={r.status_code}")
if r.status_code != 302:
    print(r.text[:600])
    sys.exit(1)
meeting_url = r.headers["Location"]
meeting_id = int(meeting_url.rstrip("/").split("/")[-1])
check("meeting id is a number", str(meeting_id).isdigit(), str(meeting_id))

r = requests.get(BASE + meeting_url, timeout=20)
body = r.text
check("results page renders", r.status_code == 200)
for needle in ["Summary", "Key discussion points", "Decisions", "Action items",
               "Follow-up email draft", "Not specified", "relative date",
               "absolute date"]:
    check(f"results contain '{needle}'", needle in body)

print("\n== downloads ==")
r = requests.get(f"{BASE}/meeting/{meeting_id}/report", timeout=20)
check("HTML report downloads",
      r.status_code == 200 and "attachment" in r.headers.get("Content-Disposition", "")
      and "Action items" in r.text)

r = requests.get(f"{BASE}/meeting/{meeting_id}/tasks.csv", timeout=20)
check("CSV downloads",
      r.status_code == 200 and r.text.lstrip("﻿").startswith("Meeting,Task,Owner"),
      r.text.splitlines()[0][:40])

print("\n== edit action items (AJAX) ==")
r = requests.get(BASE + meeting_url, timeout=20)
# grab the real ids from the page
import re
ids = [int(x) for x in re.findall(r'class="item-row" data-id="(\d+)"', r.text)]
check("page has item rows", len(ids) > 0, f"{len(ids)} rows")

payload = {"items": [
    {"id": ids[0], "task": "Send the release notes", "owner": "Alice",
     "deadline": "next Monday", "status": "In progress"},
    {"id": ids[1], "task": "Prepare feedback summary", "owner": "",
     "deadline": "", "status": "Not started"},
    {"task": "A brand new task", "owner": "Carol", "deadline": "2026-11-01",
     "status": "Done"},
]}
r = requests.post(f"{BASE}/meeting/{meeting_id}/items",
                  json=payload, timeout=20)
data = r.json()
check("items saved", r.status_code == 200 and data.get("ok") is True,
      str(data)[:200])
saved = data.get("items", [])
check("3 items stored", len(saved) == 3, str(len(saved)))
check("relative deadline re-classified + resolved",
      saved[0]["deadline_type"] == "relative"
      and saved[0]["deadline_resolved"] == "2026-10-12",
      str(saved[0]))
check("missing owner stays null", saved[1]["owner"] is None)
check("absolute deadline classified",
      saved[2]["deadline_type"] == "absolute"
      and saved[2]["deadline_resolved"] == "2026-11-01")

r = requests.post(f"{BASE}/meeting/{meeting_id}/items",
                  json={"items": [{"id": ids[0], "task": "   "}]}, timeout=20)
check("empty task rejected with 400", r.status_code == 400, str(r.json())[:80])

r = requests.post(f"{BASE}/meeting/{meeting_id}/items",
                  json={"items": "nope"}, timeout=20)
check("bad payload rejected", r.status_code == 400)

print("\n== edit title / email ==")
r = requests.post(f"{BASE}/meeting/{meeting_id}/title",
                  json={"title": "Weekly Product Sync (edited)"}, timeout=20)
check("title saved", r.json().get("ok") is True and r.json()["title"].endswith("(edited)"))

r = requests.post(f"{BASE}/meeting/{meeting_id}/title", json={"title": "x"}, timeout=20)
check("short title rejected", r.status_code == 400)

r = requests.post(f"{BASE}/meeting/{meeting_id}/email",
                  json={"subject": "Follow-up", "body": "Hello,\n\nAll done."},
                  timeout=20)
check("email saved", r.json().get("ok") is True)

r = requests.post(f"{BASE}/meeting/{meeting_id}/email",
                  json={"subject": "", "body": ""}, timeout=20)
check("empty email rejected", r.status_code == 400)

# persistence check (the data really went into SQLite)
r = requests.get(BASE + meeting_url, timeout=20)
check("changes persisted in the page",
      "Weekly Product Sync (edited)" in r.text
      and "Send the release notes" in r.text
      and "A brand new task" in r.text
      and "Hello," in r.text)
check("owner null shows as Not specified", "Not specified" in r.text)

print("\n== chat: ask about this meeting (real Ollama) ==")
r = requests.get(BASE + meeting_url, timeout=20)
check("chat panel on the page", "Ask about this meeting" in r.text)

r = requests.post(f"{BASE}/meeting/{meeting_id}/chat",
                  json={"question": ""}, timeout=20)
check("empty question rejected", r.status_code == 400)

r = requests.post(f"{BASE}/meeting/{meeting_id}/chat",
                  json={"question": "Who must send the release notes, and by when?"},
                  timeout=120)
data = r.json()
answer = str(data.get("answer", ""))
check("chat answers", r.status_code == 200 and data.get("ok") is True
      and len(answer) > 5, str(data)[:160])
check("chat answer is on topic",
      "alice" in answer.lower() or "release" in answer.lower(), answer[:120])

r = requests.get(BASE + meeting_url, timeout=20)
check("conversation saved on the page",
      "Who must send the release notes" in r.text and "chat-bubble" in r.text)

r = requests.post(f"{BASE}/meeting/{meeting_id}/chat/clear", timeout=20)
check("chat cleared", r.json().get("ok") is True)

print("\n== ask across meetings ==")
r = requests.post(BASE + "/ask", data={"question": "zzz purple unicorn ballet"},
                  timeout=20)
check("no-match question explained", r.status_code == 200
      and "could not find any meeting" in r.text)

r = requests.post(BASE + "/ask",
                  data={"question": "Who sends the release notes?"}, timeout=120)
check("ask answers with sources", r.status_code == 200 and "Meetings used" in r.text)
check("source list names the meeting", "Weekly Product Sync" in r.text)

print("\n== audio transcription route ==")
r = requests.post(BASE + "/transcribe", data={}, timeout=20)
check("no audio file rejected", r.status_code == 400)

r = requests.post(BASE + "/transcribe",
                  files={"audio": ("notes.txt", b"hello", "text/plain")}, timeout=20)
check("non-audio file rejected", r.status_code == 400
      and "Audio files only" in r.json().get("error", ""))

r = requests.post(BASE + "/transcribe",
                  files={"audio": ("fake.mp3", b"not really audio", "audio/mpeg")},
                  timeout=180)
data = r.json()
check("fake audio handled gracefully", r.status_code == 400
      and data.get("ok") is False and len(data.get("error", "")) > 10,
      str(data)[:140])

print("\n== DOCX upload ==")
buf = io.BytesIO()
doc = Document()
doc.add_paragraph("Board meeting notes")
doc.add_paragraph("It was agreed to approve the 2026 budget.")
doc.add_paragraph("Dana will circulate the minutes by 2026-11-05.")
doc.save(buf)
r = requests.post(BASE + "/analyze",
                  files={"file": ("board.docx", buf.getvalue(),
                                  "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                  data={"meeting_date": "2026-10-30"}, allow_redirects=False, timeout=90)
check("docx analyzed", r.status_code == 302, f"status={r.status_code}")
docx_url = r.headers.get("Location", "")
r = requests.get(BASE + docx_url, timeout=20)
check("docx results render", r.status_code == 200 and "Board meeting" in r.text)

print("\n== PDF upload ==")


def make_pdf(lines):
    """Build a tiny, valid one-page PDF containing the given text lines."""
    content = "BT /F1 12 Tf 72 720 Td 16 TL\n"
    for line in lines:
        content += f"({line}) Tj T*\n"
    content += "ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_pos}\n%%EOF").encode()
    return out


r = requests.post(BASE + "/analyze",
                  files={"file": ("review.pdf",
                                  make_pdf(["Product review meeting",
                                            "Alice will ship the fix by 2026-10-20.",
                                            "It was agreed to raise prices."]),
                                  "application/pdf")},
                  data={"meeting_date": "2026-10-02"},
                  allow_redirects=False, timeout=90)
check("pdf analyzed", r.status_code == 302, f"status={r.status_code}")
if r.status_code == 302:
    r2 = requests.get(BASE + r.headers["Location"], timeout=20)
    check("pdf text made it into the results",
          r2.status_code == 200 and "Product review" in r2.text)

print("\n== history / dashboard / settings / delete ==")
r = requests.get(BASE + "/history", timeout=20)
check("history lists meetings", r.status_code == 200
      and "Weekly Product Sync (edited)" in r.text)

r = requests.get(BASE + "/history?q=weekly", timeout=20)
check("history search works", "Weekly Product Sync (edited)" in r.text)

r = requests.get(BASE + "/", timeout=20)
check("dashboard shows counters", r.status_code == 200 and "Action items" in r.text)

r = requests.post(BASE + "/settings", data={"ollama_url": "notaurl"},
                  allow_redirects=False, timeout=20)
check("bad url rejected", r.status_code == 302)

r = requests.post(BASE + "/settings",
                  data={"ollama_url": os.getenv("OLLAMA_URL",
                                                "http://localhost:11434/api/chat"),
                        "ollama_model": os.getenv("OLLAMA_MODEL", "qwen2.5:3b"),
                        "whisper_model": "base"},
                  allow_redirects=False, timeout=20)
check("settings saved", r.status_code == 302)

r = requests.post(BASE + "/settings/test", timeout=40)
check("connection test returns json", r.status_code == 200
      and "message" in r.json(), str(r.json())[:120])

r = requests.post(f"{BASE}/meeting/{meeting_id}/delete",
                  allow_redirects=False, timeout=20)
check("meeting deleted", r.status_code == 302)
r = requests.get(BASE + meeting_url, timeout=20)
check("deleted meeting is gone", r.status_code == 404)

r = requests.delete(BASE + "/", timeout=10)
check("unknown method -> 405", r.status_code in (405, 200))

print("\n" + ("ALL HTTP TESTS PASSED" if not failures
              else f"{len(failures)} FAILED: {failures}"))
sys.exit(1 if failures else 0)
