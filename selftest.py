"""Quick self-test of the project's core logic (not part of the app)."""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import ai_processor as ai
import file_extractor as fx
import database as db

db.init_db()          # the app does this on startup; do it here too

failures = []


def check(label, condition, extra=""):
    print(("  OK   " if condition else "  FAIL ") + label + (f"  {extra}" if extra else ""))
    if not condition:
        failures.append(label)


print("== classify_deadline ==")
cases = [
    ("2026-10-20", "2026-10-09", "absolute", "2026-10-20"),
    ("Oct 20, 2026", "2026-10-09", "absolute", "2026-10-20"),
    ("20 October 2026", "2026-10-09", "absolute", "2026-10-20"),
    ("next Monday", "2026-10-09", "relative", "2026-10-12"),   # Friday -> next Monday
    ("next Monday", None, "relative", None),
    ("tomorrow", "2026-10-09", "relative", "2026-10-10"),
    ("today", "2026-10-09", "relative", "2026-10-09"),
    ("in 3 days", "2026-10-09", "relative", "2026-10-12"),
    ("in 2 weeks", "2026-10-09", "relative", "2026-10-23"),
    ("end of next week", "2026-10-09", "relative", "2026-10-16"),
    ("end of week", "2026-10-09", "relative", "2026-10-09"),
    ("next month", "2026-10-09", "relative", "2026-11-08"),
    ("ASAP", None, "relative", None),
    ("whenever possible", None, "unclassified", None),
    ("", None, None, None),
    ("N/A", None, None, None),
    (None, None, None, None),
]
for text, meeting_date, want_type, want_resolved in cases:
    dtype, resolved = ai.classify_deadline(text, meeting_date)
    ok = dtype == want_type
    if want_resolved:
        ok = ok and resolved == want_resolved
    if want_type == "relative" and meeting_date is None:
        ok = ok and resolved is None          # never resolve without a date
    check(f"deadline {text!r} date={meeting_date} -> {dtype}/{resolved}", ok)

print("\n== parse_json (malformed AI answers) ==")
good = '{"title":"T","summary":"S","key_points":[],"decisions":[],"action_items":[],"follow_up_email":{"subject":"s","body":"b"}}'
check("plain json", ai.parse_json(good)["title"] == "T")
check("fenced json", ai.parse_json("```json\n" + good + "\n```")["title"] == "T")
check("json with prose around it",
      ai.parse_json("Sure! Here is the result:\n" + good + "\nHope that helps.")["title"] == "T")
for bad in ("", "I could not help with that.", "{broken json", "[]"):
    try:
        ai.parse_json(bad)
        check(f"reject {bad!r}", False, "no error raised")
    except ai.AIError:
        check(f"reject {bad!r}", True)

print("\n== normalize (no invented data) ==")
result = ai.normalize({
    "title": "  Planning   call ",
    "summary": "We talked about the launch.",
    "key_points": "line one\nline two",
    "decisions": [{"decision": "Ship on the 20th", "status": "CONFIRMED"},
                  {"decision": "Maybe add dark mode", "status": "suggested"},
                  {"decision": "Rent bigger office"}],
    "action_items": [
        {"task": "Send the report", "owner": "Alice", "deadline": "next Monday"},
        {"task": "Book the room", "owner": "N/A", "deadline": None},
        {"task": "   ", "owner": "Ghost"},
        "Just a string task",
    ],
    "follow_up_email": {},
}, "Planning call notes", "2026-10-09", "test-model")

check("title trimmed", result["title"] == "Planning call")
check("summary kept", "launch" in result["summary"])
check("string key_points split", result["key_points"] == ["line one", "line two"])
check("decision statuses",
      [d["status"] for d in result["decisions"]] == ["confirmed", "suggested", "not_specified"])
check("blank task dropped", len(result["action_items"]) == 3)
check("owner kept", result["action_items"][0]["owner"] == "Alice")
check("missing owner -> None", result["action_items"][1]["owner"] is None)
check("missing deadline -> None", result["action_items"][1]["deadline"] is None)
check("relative deadline detected",
      result["action_items"][0]["deadline_type"] == "relative")
check("relative deadline resolved (meeting date known)",
      result["action_items"][0]["deadline_resolved"] == "2026-10-12",
      str(result["action_items"][0]["deadline_resolved"]))
check("status not invented", all(i["status"] == "Not started" for i in result["action_items"]))
check("fallback email generated", "Subject" not in result["email_subject"]
      and len(result["email_body"]) > 40)

print("\n== local_analyze (offline mode, no model configured) ==")
with open(os.path.join(os.path.dirname(__file__), "samples", "sample_meeting.txt"),
          encoding="utf-8") as fh:
    sample = fh.read()
out = ai.local_analyze(sample, "2026-10-09")
check("has a title", bool(out["title"]))
check("has a summary", len(out["summary"]) > 20)
check("found key points", len(out["key_points"]) > 0, str(len(out["key_points"])))
check("found decisions", len(out["decisions"]) > 0, str(len(out["decisions"])))
check("found action items", len(out["action_items"]) > 0, str(len(out["action_items"])))
check("found at least one owner", any(i["owner"] for i in out["action_items"]))
check("resolved a relative deadline", any(i["deadline_resolved"] for i in out["action_items"]))
check("email written", len(out["email_body"]) > 60)

print("\n== analyze() with no model -> offline path ==")
_env_model = os.environ.get("OLLAMA_MODEL")
os.environ["OLLAMA_MODEL"] = ""              # pretend no model is configured
db.set_setting("ollama_model", "")
try:
    check("is_configured() False in offline mode", not ai.is_configured())
    res = ai.analyze("Alice will send the report by next Friday.", "2026-10-09")
    check("offline model label", "Offline" in res["ai_model"], res["ai_model"])
finally:
    os.environ["OLLAMA_MODEL"] = _env_model or ""
    db.set_setting("ollama_model", "")        # back to the .env default

print("\n== ai_service: local Ollama via /api/chat (mock server) ==")
import ai_service as svc
from mock_ollama import MockOllama

mock = MockOllama(model="qwen2.5:3b").start()
try:
    out = svc.analyze_meeting(sample, "2026-10-09", model="qwen2.5:3b",
                              url=mock.chat_url, timeout=(5, 20))
    check("returns a meeting_summary", len(out["meeting_summary"]) >= 20)
    check("returns key_points", bool(out["key_points"]))
    check("decisions carry a status", bool(out["decisions"]) and all(
        d["status"] in ("confirmed", "suggested", "not_specified")
        for d in out["decisions"]))
    items = out["action_items"]
    check("items carry task/owner/deadline/status", bool(items) and all(
        {"task", "owner", "deadline", "status"} <= set(i) for i in items))
    check("owners only from the transcript", all(
        i["owner"] in (None, "Alice", "Bob", "Carol", "Dave") for i in items),
        str(sorted({i["owner"] for i in items if i["owner"]})))
    check("a missing deadline stays null", any(i["deadline"] is None for i in items))
    check("task-embedded deadline recovered (relative)",
          ai.extract_task_deadline("Send the release notes by next Monday")
          == "next Monday")
    check("task-embedded deadline recovered (absolute)",
          ai.extract_task_deadline("Book the room - due 2026-10-16")
          == "2026-10-16")
    check("task without a deadline stays None", ai.extract_task_deadline(
        "Prepare the customer feedback summary (no deadline given)") is None)
    check("unrecognized wording is not guessed",
          ai.extract_task_deadline("Do the thing by whenever possible") is None)
    check("status limited to the allowed values",
          all(i["status"] in svc.TASK_STATUSES for i in items))
    check("follow-up email has subject + body",
          bool(out["follow_up_email"]["subject"])
          and len(out["follow_up_email"]["body"]) > 40)

    req = mock.last_request
    check("POST /api/chat used the right model + no streaming",
          req.get("model") == "qwen2.5:3b" and req.get("stream") is False)
    check("transcript reached the model", "Alice" in mock.transcript)
    check("system prompt demands ONE JSON object",
          "JSON" in (req.get("messages") or [{}])[0].get("content", ""))

    # --- validation rules (never invent anything) -------------------------
    validated = svc.validate_ai_response({
        "title": "T",
        "meeting_summary": "A summary that is definitely long enough to keep.",
        "key_points": ["p1"],
        "decisions": [{"decision": "d1", "status": "weird"}],
        "action_items": [
            {"task": "T1", "owner": None, "deadline": None, "status": "Bogus"},
            {"task": "T2", "owner": "N/A", "deadline": "N/A", "status": "Done"},
        ],
        "follow_up_email": {"subject": "s", "body": "b"},
    })
    check("null owner stays None", validated["action_items"][0]["owner"] is None)
    check("null deadline stays None", validated["action_items"][0]["deadline"] is None)
    check("invalid status -> Not started",
          validated["action_items"][0]["status"] == "Not started")
    check("placeholder owner/deadline -> None",
          validated["action_items"][1]["owner"] is None
          and validated["action_items"][1]["deadline"] is None)
    check("valid status kept", validated["action_items"][1]["status"] == "Done")
    check("unknown decision status -> not_specified",
          validated["decisions"][0]["status"] == "not_specified")

    # --- malformed answers are reported, never crashed on -----------------
    for mode, needle in (("fenced", "meeting_summary"),
                         ("badjson", "not valid JSON"),
                         ("empty", "empty answer"),
                         ("notjsonbody", "not valid JSON"),
                         ("missing_summary", "meeting_summary"),
                         ("http404", "ollama pull"),
                         ("http500", "internal error")):
        mock.mode = mode
        try:
            got = svc.analyze_meeting(sample, "2026-10-09",
                                      url=mock.chat_url, timeout=(5, 20))
            if mode == "fenced":
                check("markdown-fenced JSON accepted",
                      len(got["meeting_summary"]) >= 20)
            else:
                check(f"handles {mode} gracefully", False, "no error raised")
        except svc.AIServiceError as exc:
            if mode == "fenced":
                check("markdown-fenced JSON accepted", False, str(exc))
            else:
                check(f"handles {mode} gracefully",
                      needle.lower() in str(exc).lower(), str(exc))
    mock.mode = "good"

    # --- connection problems ----------------------------------------------
    try:
        svc.analyze_meeting(sample, model="qwen2.5:3b",
                            url="http://127.0.0.1:9/api/chat", timeout=(2, 5))
        check("unreachable Ollama reported", False, "no error raised")
    except svc.AIServiceError as exc:
        check("unreachable Ollama reported",
              "could not reach ollama" in str(exc).lower(), str(exc))

    mock.delay = 2.0
    try:
        svc.analyze_meeting(sample, model="qwen2.5:3b",
                            url=mock.chat_url, timeout=(5, 0.4))
        check("slow model -> timeout message", False, "no error raised")
    except svc.AIServiceError as exc:
        check("slow model -> timeout message", "too long" in str(exc).lower(),
              str(exc))
    finally:
        mock.delay = 0.0

    # --- the Settings-page connection test ---------------------------------
    ok, msg = svc.test_connection(model="qwen2.5:3b", url=mock.chat_url)
    check("test_connection succeeds", ok, msg)

    other = MockOllama(tags_models=["llama3.2:1b"]).start()
    try:
        ok, msg = svc.test_connection(model="qwen2.5:3b", url=other.chat_url)
        check("undownloaded model -> 'ollama pull' hint",
              (not ok) and "ollama pull" in msg, msg)
    finally:
        other.stop()

    ok, msg = svc.test_connection(model="qwen2.5:3b",
                                  url="http://127.0.0.1:9/api/chat")
    check("test_connection unreachable -> friendly message",
          (not ok) and "could not reach" in msg.lower(), msg)

    # --- end to end: config in SQLite -> ai_processor -> normalized result -
    db.set_setting("ollama_url", mock.chat_url)
    db.set_setting("ollama_model", "qwen2.5:3b")
    try:
        check("ai_processor sees the configured model", ai.is_configured())
        res2 = ai.analyze(sample, "2026-10-09")
        check("analyze() used the configured model",
              res2["ai_model"] == "qwen2.5:3b", res2["ai_model"])
        check("relative deadline classified + resolved", any(
            i["deadline_type"] == "relative" and i["deadline_resolved"] == "2026-10-12"
            for i in res2["action_items"]))
        check("absolute deadline classified",
              any(i["deadline_type"] == "absolute" for i in res2["action_items"]))
        check("missing owner still None after normalize",
              all(i["owner"] is None or isinstance(i["owner"], str)
                  for i in res2["action_items"]))
        ok, msg = ai.test_connection()
        check("ai_processor.test_connection", ok, msg)

        mock.mode = "no_deadlines"       # model buries deadlines in the task
        try:
            res3 = ai.analyze(sample, "2026-10-09")
            check("deadline buried in the task sentence is recovered", any(
                i["deadline"] == "next Monday" and i["deadline_type"] == "relative"
                and i["deadline_resolved"] == "2026-10-12"
                for i in res3["action_items"]))
        finally:
            mock.mode = "good"
    finally:
        db.set_setting("ollama_url", "")
        db.set_setting("ollama_model", "")
finally:
    mock.stop()

print("\n== real Ollama (optional - skipped when it is not running) ==")
ok, msg = svc.test_connection()               # values come from .env
if ok:
    check("real Ollama reachable", True, msg)
else:
    print("  SKIP  " + msg)


print("\n== chat + ask (Q&A over meetings) ==")
# --- saved conversations (database) ---
m_id = db.create_meeting({"title": "Chat test", "summary": "A short summary text.",
                          "source_text": "Alice will send the report by next Monday.",
                          "ai_model": "test"})
db.add_chat(m_id, "user", "Who sends the report?")
db.add_chat(m_id, "assistant", "Alice, by next Monday.")
chat = db.list_chats(m_id)
check("chat messages saved", [c["role"] for c in chat] == ["user", "assistant"])
check("chat content kept", chat[0]["content"] == "Who sends the report?")
db.clear_chats(m_id)
check("chat cleared", db.list_chats(m_id) == [])
db.add_chat(m_id, "user", "and again")
db.delete_meeting(m_id)
check("chat deleted with the meeting", db.list_chats(m_id) == [])

# --- chat_context builder ---
ctx = ai.chat_context(
    {"title": "T", "meeting_date": "2026-10-09", "summary": "S",
     "key_points": ["p1"],
     "decisions": [{"decision": "d1", "status": "confirmed"}],
     "action_items": [{"task": "t1", "owner": "Alice", "deadline": None,
                       "status": "Not started"}],
     "source_text": "raw notes here"},
    [{"role": "user", "content": "earlier question?"}])
check("chat_context has title + date", "T (2026-10-09)" in ctx)
check("chat_context has the items", "Alice" in ctx and "Not specified" in ctx)
check("chat_context has the history", "earlier question?" in ctx)

# --- keyword retrieval across meetings ---
rid = db.create_meeting({"title": "Zebra project sync",
                         "summary": "Talk about the zebra launch.",
                         "source_text": "The zebra project ships in November. Carol owns it.",
                         "ai_model": "test"})
hits = ai.find_relevant_meetings("When does the zebra project ship?")
check("relevant meeting found", any(h["id"] == rid for h in hits),
      str([h["title"] for h in hits]))
check("best match comes first", bool(hits) and hits[0]["id"] == rid)
check("unrelated question finds nothing",
      ai.find_relevant_meetings("purple unicorn ballet") == [])

# --- plain-text answers through the mock Ollama ---
mock2 = MockOllama(model="qwen2.5:3b", mode="plain").start()
try:
    answer = svc.ask_meeting("Meeting: X\nSummary: Y", "Who sends the report?",
                             model="qwen2.5:3b", url=mock2.chat_url,
                             timeout=(5, 20))
    check("ask_meeting returns plain text", "Alice" in answer, answer)
    req = mock2.last_request
    check("chat uses no JSON mode", "format" not in req)
    check("chat starts with a system prompt",
          (req.get("messages") or [{}])[0].get("role") == "system")
    answer2 = svc.ask_across("=== Meeting: A ===\nnotes", "Any news?",
                             model="qwen2.5:3b", url=mock2.chat_url,
                             timeout=(5, 20))
    check("ask_across returns plain text", bool(answer2))
    try:
        svc.ask_meeting("ctx", "x" * (svc.ASK_MAX_QUESTION + 1),
                        url=mock2.chat_url)
        check("overlong question rejected", False, "no error raised")
    except svc.AIServiceError as exc:
        check("overlong question rejected", "too long" in str(exc).lower(),
              str(exc))

    # --- end to end: SQLite settings -> ai_processor -> mock answer ---
    db.set_setting("ollama_url", mock2.chat_url)
    db.set_setting("ollama_model", "qwen2.5:3b")
    try:
        row = db.get_meeting(rid)
        reply = ai.ask_meeting(row, [], "What about the zebra project?")
        check("ai_processor.ask_meeting works", "Alice" in reply, reply)
        reply2 = ai.ask_all([row], "Any news?")
        check("ai_processor.ask_all works", bool(reply2))
    finally:
        db.set_setting("ollama_url", "")
        db.set_setting("ollama_model", "")
finally:
    mock2.stop()
db.delete_meeting(rid)

print("\n== audio_service (optional faster-whisper) ==")
import audio_service as auc
check("audio accepts mp3 (any case)", auc.validate_audio("meeting.MP3") == ".mp3")
check("default model size is valid", auc.default_model_size() in auc.MODEL_SIZES)
os.environ["WHISPER_MODEL"] = "tiny"
check("WHISPER_MODEL from the environment is honoured",
      auc.default_model_size() == "tiny")
os.environ["WHISPER_MODEL"] = "not-a-model"
check("invalid WHISPER_MODEL falls back safely",
      auc.default_model_size() in auc.MODEL_SIZES)
os.environ.pop("WHISPER_MODEL", None)
try:
    auc.validate_audio("evil.exe")
    check("audio rejects non-audio file", False, "no error raised")
except auc.AudioValidationError:
    check("audio rejects non-audio file", True)
print("  INFO faster-whisper installed: " +
      ("yes" if auc.is_available() else "no"))
if not auc.is_available():
    try:
        auc.transcribe(b"x", "a.mp3")
        check("missing package reported", False, "no error raised")
    except auc.AudioValidationError as exc:
        check("missing package reported", "faster-whisper" in str(exc), str(exc))
try:
    auc.transcribe(b"", "empty.mp3")
    check("empty audio rejected", False, "no error raised")
except auc.AudioValidationError:
    check("empty audio rejected", True)


print("\n== database round trip ==")
db.init_db()
mid = db.create_meeting({
    "title": "Unit test meeting", "summary": "s",
    "key_points": ["p1"], "decisions": [{"decision": "d", "status": "confirmed"}],
    "email_subject": "sub", "email_body": "body",
    "meeting_date": "2026-10-09", "source_filename": None,
    "source_text": "text", "ai_model": "test",
    "action_items": [{"task": "T1", "owner": None, "deadline": "next Monday",
                      "deadline_type": "relative", "deadline_resolved": "2026-10-12",
                      "status": "Not started"}],
})
loaded = db.get_meeting(mid)
check("meeting saved", loaded is not None and loaded["title"] == "Unit test meeting")
check("json fields parsed", loaded["key_points"] == ["p1"])
check("item saved", loaded["action_items"][0]["task"] == "T1")

saved = db.save_action_items(mid, [
    {"id": loaded["action_items"][0]["id"], "task": "T1 edited", "owner": "Bob",
     "deadline": "2026-11-01", "deadline_type": "absolute",
     "deadline_resolved": "2026-11-01", "status": "Done"},
    {"task": "T2 new", "owner": None, "deadline": None,
     "deadline_type": None, "deadline_resolved": None, "status": "Not started"},
])
check("item updated + inserted", len(saved) == 2 and saved[0]["task"] == "T1 edited")
check("status updated", saved[0]["status"] == "Done")

saved2 = db.save_action_items(mid, [saved[1] | {"id": saved[1]["id"]}])
check("item deleted when omitted", len(saved2) == 1 and saved2[0]["task"] == "T2 new")
check("search works", any(m["id"] == mid for m in db.list_meetings("Unit test")))
stats = db.dashboard_stats()
check("dashboard stats", stats["totals"]["meetings"] >= 1)
db.delete_meeting(mid)
check("cascade delete", db.get_meeting(mid) is None)

print("\n== file extraction ==")
try:
    fx.validate_filename("evil.exe")
    check("reject .exe", False)
except fx.FileValidationError:
    check("reject .exe", True)
check("accept .PDF uppercase", fx.extension_of("notes.PDF") == "pdf")

try:
    fx.extract_text(b"", "empty.txt")
    check("empty txt rejected", False)
except fx.FileValidationError as exc:
    check("empty txt rejected", "empty" in str(exc).lower(), str(exc))

try:
    fx.extract_text(b"%PDF-1.4 not really", "fake.pdf")
    check("unreadable pdf rejected", False)
except fx.FileValidationError:
    check("unreadable pdf rejected", True)

txt = fx.extract_text("Hello\r\n\r\n\r\n\r\nWorld\x00!".encode("utf-16"), "a.txt")
check("txt decoded + tidied", "Hello" in txt and "\x00" not in txt and "\n\n\n" not in txt)

# build a real .docx
from docx import Document
import io
buf = io.BytesIO()
doc = Document()
doc.add_paragraph("Team standup notes")
doc.add_paragraph("Bob will update the roadmap by next Tuesday.")
t = doc.add_table(rows=1, cols=2)
t.rows[0].cells[0].text = "Owner"
t.rows[0].cells[1].text = "Alice"
buf.seek(0)
out_docx = buf.getvalue()
buf2 = io.BytesIO(); doc.save(buf2); out_docx = buf2.getvalue()
text_docx = fx.extract_text(out_docx, "notes.docx")
check("docx paragraphs", "standup" in text_docx)
check("docx tables", "Alice" in text_docx)

# build a real .pdf with text
def make_pdf(lines):
    content = "BT /F1 12 Tf 72 720 Td 16 TL\n"
    for line in lines:
        content += f"({line}) Tj T*\n"
    content += "ET"
    objs = []
    objs.append("<< /Type /Catalog /Pages 2 0 R >>")
    objs.append("<< /Type /Pages /Kids [3 0 R] /Count 1 >>")
    objs.append("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                "/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>")
    objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objs.append(f"<< /Length {len(content)} >>\nstream\n{content}\nendstream")
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref_pos = len(out)
    out += f"xref\n0 {len(objs)+1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objs)+1} /Root 1 0 R >>\n"
            f"startxref\n{xref_pos}\n%%EOF").encode()
    return out

pdf = make_pdf(["Product review meeting", "Alice will ship the fix by 2026-10-20.",
                "It was agreed to raise prices next month."])
text_pdf = fx.extract_text(pdf, "meeting.pdf")
check("pdf text extracted", "Product review" in text_pdf, repr(text_pdf[:60]))

print("\n" + ("ALL TESTS PASSED" if not failures else f"{len(failures)} FAILED: {failures}"))
sys.exit(1 if failures else 0)
